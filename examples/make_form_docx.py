# -*- coding: utf-8 -*-
"""把 FORMULA-AGH_填表文本.md 转成 Word 文档。

为什么由 md 生成而不是手写 docx：填表文本是唯一事实来源，
手写一份 docx 迟早与它漂移（本项目已多次吃过这个亏）。

    python examples/make_form_docx.py
"""
import os
import re
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
SRC = os.path.join(os.path.dirname(ROOT), "FORMULA-AGH_填表文本.md")
OUT = os.path.join(os.path.dirname(ROOT), "FORMULA-AGH_填表文本.docx")
# 可用 --source / --out 覆盖（默认仍是填表文本）

CJK = "宋体"
MONO = "Consolas"


def set_cjk(run, name=CJK, ascii_only=False):
    """设置字体。ascii_only=True 时只把 ASCII 设为等宽，中文仍用宋体——
    Consolas 没有中文字形，把 eastAsia 也设成它会掉进回退字体，排版不可控。"""
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    fonts.set(qn("w:eastAsia"), CJK if ascii_only else name)
    fonts.set(qn("w:ascii"), name)
    fonts.set(qn("w:hAnsi"), name)


def shade(cell, fill="F2F4F7"):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def add_code_block(doc, text):
    """命令类短块原样输出；长文（含标题/加粗）按 Markdown 渲染。

    填表文本里，作品详情也放在围栏里（便于整段复制），但它其实是 Markdown 正文，
    原样输出会把 ## 和 ** 直接印在纸上。这里按长度与是否含标题判断，分两种处理。
    """
    lines = text.split("\n")
    longform = len(text) > 300 or any(l.startswith("#") for l in lines)
    if longform:
        render_markdown(doc, lines)
        return
    if len(text) > 160 and not text.lstrip().startswith(("python", ".\\", "powershell", "cd ", "set ")):
        # 长中文纯文本块：按正文段落排版，等宽字体对中文又难读又容易溢出
        render_prose(doc, text)
        return
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    table.autofit = True
    cell = table.cell(0, 0)
    shade(cell)
    cell.text = ""
    for n, line in enumerate(lines):
        para = cell.paragraphs[0] if n == 0 else cell.add_paragraph()
        para.paragraph_format.space_after = Pt(2)
        run = para.add_run(line)
        set_cjk(run, MONO, ascii_only=True)
        run.font.size = Pt(9.5)
    doc.add_paragraph()


def render_prose(doc, text):
    """把一段长纯文本按空行拆成正文段落。"""
    for chunk in text.split("\n\n"):
        chunk = chunk.strip()
        if not chunk:
            continue
        for line in chunk.split("\n"):
            para = doc.add_paragraph()
            set_cjk(para.add_run(line.strip()))


def render_markdown(doc, lines):
    """把一段 Markdown 渲染进文档（供长文使用）。"""
    i = 0
    pending = []
    while i < len(lines):
        line = lines[i].rstrip()
        if line.startswith("|"):
            pending.append(line); i += 1; continue
        if pending:
            add_md_table(doc, pending); pending = []
        if not line.strip():
            i += 1; continue
        m = re.match(r"^(#{1,6})\s+(.*)$", line)
        if m:
            lvl = min(len(m.group(1)), 4)
            head = doc.add_heading("", level=lvl)
            run = head.add_run(re.sub(r"\*\*|`", "", m.group(2)))
            set_cjk(run)
            i += 1; continue
        txt = line
        indented = txt.startswith("    ")
        if txt.startswith("- [ ] "):
            txt = "\u2610 " + txt[6:]
        elif txt.startswith("- "):
            txt = "\u2022 " + txt[2:]
        txt = re.sub(r"\*\*(.+?)\*\*", r"\1", txt)
        txt = re.sub(r"`(.+?)`", r"\1", txt)
        para = doc.add_paragraph()
        if indented or txt.startswith(("\u2022", "\u2610")):
            para.paragraph_format.left_indent = Cm(0.6)
        set_cjk(para.add_run(txt.strip()))
        i += 1
    if pending:
        add_md_table(doc, pending)

def add_md_table(doc, rows):
    header = [c.strip() for c in rows[0].strip("|").split("|")]
    body = [[c.strip() for c in r.strip("|").split("|")] for r in rows[2:]]
    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Light Grid Accent 1"
    for i, h in enumerate(header):
        cell = table.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(re.sub(r"\*\*|`", "", h))
        set_cjk(run)
        run.bold = True
        run.font.size = Pt(10)
    for row in body:
        cells = table.add_row().cells
        for i, v in enumerate(row[:len(header)]):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(re.sub(r"\*\*|`", "", v))
            set_cjk(run)
            run.font.size = Pt(10)
    doc.add_paragraph()


def build(lines):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin = sec.bottom_margin = Cm(2.0)
    style = doc.styles["Normal"]
    style.font.size = Pt(10.5)
    style.font.name = CJK
    style.element.rPr.rFonts.set(qn("w:eastAsia"), CJK)

    i = 0
    pending_table = []
    while i < len(lines):
        line = lines[i].rstrip()
        if line.startswith("```"):
            block, i = [], i + 1
            while i < len(lines) and not lines[i].startswith("```"):
                block.append(lines[i]); i += 1
            add_code_block(doc, "\n".join(block))
            i += 1
            continue
        if line.startswith("|"):
            pending_table.append(line); i += 1
            continue
        if pending_table:
            add_md_table(doc, pending_table); pending_table = []
        if not line.strip():
            i += 1; continue
        if line.startswith("---"):
            i += 1; continue
        m = re.match(r"^(#{1,4})\s+(.*)$", line)
        if m:
            level = len(m.group(1))
            text = re.sub(r"\*\*|`", "", m.group(2))
            if level == 1:
                p = doc.add_heading("", level=0)
            else:
                p = doc.add_heading("", level=min(level - 1, 4))
            run = p.add_run(text)
            set_cjk(run)
            i += 1; continue
        if line.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.8)
            run = p.add_run(re.sub(r"\*\*|`", "", line[2:]))
            set_cjk(run)
            run.italic = True
            run.font.size = Pt(10)
            i += 1; continue
        text = line
        bullet = False
        if text.startswith("- [ ] "):
            text = "\u2610 " + text[6:]; bullet = True
        elif text.startswith("- "):
            text = "\u2022 " + text[2:]; bullet = True
        text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
        text = re.sub(r"`(.+?)`", r"\1", text)
        p = doc.add_paragraph()
        if bullet:
            p.paragraph_format.left_indent = Cm(0.6)
        elif text.startswith("    "):
            p.paragraph_format.left_indent = Cm(0.8)
        run = p.add_run(text.strip())
        set_cjk(run)
        i += 1
    if pending_table:
        add_md_table(doc, pending_table)
    doc.save(OUT)
    return OUT


def main():
    import argparse
    global SRC, OUT
    ap = argparse.ArgumentParser()
    ap.add_argument("--source", default=None)
    ap.add_argument("--out", default=None)
    args = ap.parse_args()
    if args.source:
        SRC = args.source
    if args.out:
        OUT = args.out
    if not os.path.isfile(SRC):
        print("未找到源文件：%s" % SRC)
        return 0
    return _run()


def _run():
    # 源文件是本队的提交表填写文本，属于提交材料而不是作品本身，因此放在仓库外。
    # 仓库里没有它时不要报栈，给出可读提示即可。
    if not os.path.isfile(SRC):
        print("未找到源文件，跳过（它属于提交材料，不在仓库内）：")
        print("   ", SRC)
        return 0
    with open(SRC, encoding="utf-8") as fh:
        lines = fh.read().split("\n")
    path = build(lines)
    print("已写出", path)
    return 0


if __name__ == "__main__":
    sys.exit(main())