# -*- coding: utf-8 -*-
"""把用户已有的数据文件转换成项目要求的任务格式。

项目要求的输入是 tasks/<task_id>/{data.csv, meta.json}：
data.csv 是一份带表头的数值表，目标列固定叫 y；meta.json 声明变量、单位、
自由参数、采样区间与可追溯出处。

真实用户的数据往往不是这个形状：列名任意、目标量混在中间、可能是 TSV / JSON /
Excel、带空行或空单元格。本模块把这些输入规范化后落盘，并**立即用契约校验器
验一遍**，让问题在导入阶段暴露，而不是等到验证阶段才发现。

支持格式：
  .csv / .tsv / .txt   分隔符自动嗅探（逗号、分号、制表符、竖线）
  .json                记录数组 [{列: 值}, ...] 或列对象 {列: [值, ...]}
  .xlsx / .xlsm        Excel 首个工作表（需要 openpyxl，缺失时给出明确提示）
"""

import csv
import json
import os
from datetime import datetime, timezone
from typing import Dict, List, Mapping, Optional, Sequence, Tuple

from .contract import validate_task

SUPPORTED_EXT = (".csv", ".tsv", ".txt", ".json", ".xlsx", ".xlsm")
TARGET_COLUMN = "y"


class DataImportError(Exception):
    """输入无法转换成合法任务时抛出，消息面向使用者。"""


# ---------------------------------------------------------------------------
# 读取
# ---------------------------------------------------------------------------

def _read_delimited(path: str) -> Tuple[List[List[str]], Dict[str, object]]:
    with open(path, encoding="utf-8-sig", newline="") as fh:
        sample = fh.read(8192)
        fh.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
            delimiter = dialect.delimiter
        except csv.Error:
            delimiter = "\t" if path.lower().endswith(".tsv") else ","
        rows = [list(r) for r in csv.reader(fh, delimiter=delimiter)]
    return rows, {"format": "delimited", "delimiter": delimiter}


def _read_json(path: str) -> Tuple[List[List[str]], Dict[str, object]]:
    with open(path, encoding="utf-8") as fh:
        obj = json.load(fh)

    if isinstance(obj, list):
        if not obj or not all(isinstance(r, Mapping) for r in obj):
            raise DataImportError(
                "JSON 记录数组必须是非空的 {列名: 值} 对象数组，例如 [{\"a\":1,\"y\":2}]")
        header: List[str] = []
        for rec in obj:
            for key in rec:
                if key not in header:
                    header.append(str(key))
        rows = [[rec.get(h, "") for h in header] for rec in obj]
        return [header] + [[_fmt_cell(c) for c in r] for r in rows], {"format": "json-records"}

    if isinstance(obj, Mapping):
        cols = {str(k): v for k, v in obj.items()}
        if cols and all(isinstance(v, list) for v in cols.values()):
            header = list(cols.keys())
            length = max(len(v) for v in cols.values())
            rows = [[cols[h][i] if i < len(cols[h]) else "" for h in header]
                    for i in range(length)]
            return [header] + [[_fmt_cell(c) for c in r] for r in rows], {"format": "json-columns"}
        raise DataImportError(
            "JSON 对象必须是 {列名: [值, ...]} 的列式结构，或改用记录数组")

    raise DataImportError("JSON 顶层必须是数组或对象，实际是 " + type(obj).__name__)


def _fmt_cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "1" if value else "0"
    return str(value)


def _read_xlsx(path: str) -> Tuple[List[List[str]], Dict[str, object]]:
    try:
        from openpyxl import load_workbook  # type: ignore
    except ImportError:
        raise DataImportError(
            "读取 Excel 需要 openpyxl（pip install openpyxl）；"
            "或在 Excel 里另存为 CSV 后再导入")
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        rows = [[_fmt_cell(c) for c in row] for row in ws.iter_rows(values_only=True)]
    finally:
        wb.close()
    return rows, {"format": "xlsx", "sheet": wb.sheetnames[0]}


def read_table(path: str) -> Tuple[List[str], List[List[str]], Dict[str, object]]:
    """读取任意受支持的表格文件，返回 (表头, 数据行, 读取信息)。"""
    if not os.path.isfile(path):
        raise DataImportError("找不到输入文件: " + path)
    ext = os.path.splitext(path)[1].lower()
    if ext not in SUPPORTED_EXT:
        raise DataImportError(
            "不支持的扩展名 " + ext + "；支持 " + " / ".join(SUPPORTED_EXT))

    if ext == ".json":
        rows, info = _read_json(path)
    elif ext in (".xlsx", ".xlsm"):
        rows, info = _read_xlsx(path)
    else:
        rows, info = _read_delimited(path)

    header, body = _normalize(rows, path)
    info["raw_rows"] = len(body)
    return header, body, info


def _normalize(rows: Sequence[Sequence[str]], path: str) -> Tuple[List[str], List[List[str]]]:
    """去掉全空行、定位表头、裁剪单元格空白。"""
    cleaned: List[List[str]] = []
    for row in rows:
        cells = ["" if c is None else str(c).strip() for c in row]
        if any(cells):
            cleaned.append(cells)
    if not cleaned:
        raise DataImportError("输入文件没有有效内容: " + path)

    header = cleaned[0]
    if not any(header):
        raise DataImportError("第一行表头为空")
    dup = sorted({h for h in header if header.count(h) > 1})
    if dup:
        raise DataImportError("表头存在重复列名: " + ", ".join(dup))
    if any(h == "" for h in header):
        raise DataImportError("表头存在空列名；请给每一列一个名字")

    body = [r for r in cleaned[1:] if any(c for c in r)]
    if not body:
        raise DataImportError("只有表头、没有数据行")
    return header, body


# ---------------------------------------------------------------------------
# 规范化
# ---------------------------------------------------------------------------

def _to_float_rows(header: Sequence[str], body: Sequence[Sequence[str]]) -> List[List[float]]:
    """整表转 float；任何一列出现非数值都报错，并指出具体位置。"""
    width = len(header)
    out: List[List[float]] = []
    bad: List[str] = []
    for lineno, row in enumerate(body, start=2):
        if len(row) != width:
            if len(bad) < 5:
                bad.append("第 %d 行列数 %d，表头 %d 列" % (lineno, len(row), width))
            continue
        values: List[float] = []
        for col, cell in zip(header, row):
            text = "" if cell is None else str(cell).strip()
            if text == "":
                if len(bad) < 5:
                    bad.append("第 %d 行 %s 列为空" % (lineno, col))
                values = []
                break
            try:
                number = float(text)
            except ValueError:
                if len(bad) < 5:
                    bad.append("第 %d 行 %s 列不是数值: %r" % (lineno, col, text[:24]))
                values = []
                break
            if number != number or number in (float("inf"), float("-inf")):
                if len(bad) < 5:
                    bad.append("第 %d 行 %s 列是 NaN/Inf" % (lineno, col))
                values = []
                break
            values.append(number)
        if values:
            out.append(values)
    if bad:
        raise DataImportError(
            "数据含无法转换的单元格（契约要求整表数值、不允许空值）:\n  - "
            + "\n  - ".join(bad))
    return out


def _sampling_ranges(header: Sequence[str], var_names: Sequence[str],
                     rows: Sequence[Sequence[float]]) -> Dict[str, List[float]]:
    index = {name: i for i, name in enumerate(header)}
    ranges: Dict[str, List[float]] = {}
    for name in var_names:
        col = [r[index[name]] for r in rows]
        lo, hi = min(col), max(col)
        if lo == hi:
            raise DataImportError(
                "自变量 " + name + " 在数据中是常量（min == max），无法构成发现任务")
        ranges[name] = [lo, hi]
    return ranges


def _write_csv(path: str, header: Sequence[str], rows: Sequence[Sequence[float]]) -> None:
    with open(path, "w", encoding="utf-8", newline="\n") as fh:
        writer = csv.writer(fh, lineterminator="\n")
        writer.writerow(header)
        for row in rows:
            writer.writerow([repr(float(v)) if v != int(v) else str(int(v)) for v in row])


# ---------------------------------------------------------------------------
# 主流程
# ---------------------------------------------------------------------------

def convert(source: str,
            task_id: str,
            target: str = TARGET_COLUMN,
            variables: Optional[Sequence[str]] = None,
            layer: str = "base",
            source_url: str = "",
            units: Optional[Mapping[str, str]] = None,
            free_parameters: Optional[Sequence[str]] = None,
            domain: str = "",
            phenomenon: str = "",
            tasks_root: str = "tasks",
            reference_root: str = "reference",
            formula: str = "",
            force: bool = False) -> Dict[str, object]:
    """把一个数据文件转换成 tasks/<task_id>/{data.csv, meta.json}。

    目标列会被重命名为 y（原列名记录进 meta.import.target_column_original）。
    返回报告字典，其中 findings 是契约校验结果——**错误不为空表示本次导入不可用**。
    """
    task_id = (task_id or "").strip()
    if not task_id:
        raise DataImportError("task_id 不能为空")
    if any(ch in task_id for ch in '\\/:*?"<>|'):
        raise DataImportError("task_id 不能包含路径分隔符或非法字符: " + task_id)
    if layer not in ("base", "challenge"):
        raise DataImportError("layer 只能是 base 或 challenge，实际是 " + layer)

    task_dir = os.path.join(tasks_root, task_id)

    header, body, read_info = read_table(source)
    if target not in header:
        raise DataImportError(
            "目标列 " + target + " 不在表头里；实际列名: " + ", ".join(header))

    var_names = [str(v).strip() for v in variables] if variables else \
        [c for c in header if c != target]
    if not var_names:
        raise DataImportError("没有自变量：除目标列外至少还需要一列")
    missing = [v for v in var_names if v not in header]
    if missing:
        raise DataImportError("自变量不在表头里: " + ", ".join(missing))
    if target in var_names:
        raise DataImportError("目标列不能同时作为自变量: " + target)
    if TARGET_COLUMN in var_names:
        raise DataImportError(
            "自变量不能叫 " + TARGET_COLUMN + "（该名字保留给目标量）")
    if len(var_names) < 2:
        raise DataImportError("自变量至少 2 个，否则不构成发现任务")

    rows = _to_float_rows(header, body)
    if len(rows) < 50:
        raise DataImportError(
            "样本量 %d 低于验证引擎硬下限 50（config 建议 >= 200）" % len(rows))

    # 目标列重命名为 y，其余保持原顺序
    out_header = ["y" if c == target else c for c in header]
    target_index = header.index(target)
    out_rows = [[r[target_index]] + [v for i, v in enumerate(r) if i != target_index]
                for r in rows]
    out_header = ["y"] + [c for c in header if c != target]

    sampling = _sampling_ranges(header, var_names, rows)

    # 覆盖检查放在输入全部通过之后：输入有问题时应先报输入的问题，
    # 而不是被"目录已存在"挡住，导致使用者以为数据没问题。
    if os.path.exists(task_dir) and not force:
        raise DataImportError(
            "任务目录已存在: " + task_dir + "；确认要覆盖请加 --force")

    os.makedirs(task_dir, exist_ok=True)
    _write_csv(os.path.join(task_dir, "data.csv"), out_header, out_rows)

    sources = [source_url] if source_url else [os.path.abspath(source)]
    meta: Dict[str, object] = {
        "task_id": task_id,
        "var_names": var_names,
        "layer": layer,
        "source": sources,
        "sampling": sampling,
        "sample_count": len(out_rows),
        "import": {
            "imported_from": os.path.abspath(source),
            "imported_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
            "read_info": read_info,
            "target_column_original": target,
            "target_column_renamed_to": TARGET_COLUMN,
            "sampling_inferred_from_data": True,
        },
    }
    if units:
        meta["units"] = {k: str(v) for k, v in units.items()}
    if free_parameters:
        meta["free_parameters"] = [str(p) for p in free_parameters]
    if domain:
        meta["domain"] = domain
    if phenomenon:
        meta["phenomenon"] = phenomenon

    with open(os.path.join(task_dir, "meta.json"), "w", encoding="utf-8") as fh:
        json.dump(meta, fh, ensure_ascii=False, indent=2)
        fh.write("\n")

    reference_path = ""
    if formula:
        os.makedirs(reference_root, exist_ok=True)
        reference_path = os.path.join(reference_root, task_id + ".json")
        with open(reference_path, "w", encoding="utf-8") as fh:
            json.dump({
                "task_id": task_id,
                "formula": formula,
                "free_parameters": [str(p) for p in (free_parameters or [])],
                "true_constants": {},
                "source": sources,
            }, fh, ensure_ascii=False, indent=2)
            fh.write("\n")

    findings = validate_task(task_dir, task_id, require_split=False)
    errors = [f for f in findings if getattr(f, "severity", "") == "error"]
    return {
        "task_id": task_id,
        "task_dir": task_dir,
        "data_csv": os.path.join(task_dir, "data.csv"),
        "meta_json": os.path.join(task_dir, "meta.json"),
        "reference_json": reference_path,
        "rows": len(out_rows),
        "var_names": var_names,
        "target_renamed_from": target,
        "sampling": sampling,
        "findings": [{"severity": f.severity, "code": f.code, "message": f.message}
                     for f in findings],
        "errors": len(errors),
        "warnings": len(findings) - len(errors),
    }
